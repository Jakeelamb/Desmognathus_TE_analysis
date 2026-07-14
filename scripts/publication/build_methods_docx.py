#!/usr/bin/env python3
"""Build the GBE working Methods DOCX from its auditable Markdown source."""

from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
METHODS = ROOT / "Publication" / "methods"
SOURCE = METHODS / "Materials_and_Methods_GBE_draft.md"
REFERENCE = METHODS / "GBE_methods_reference.docx"
OUTPUT = METHODS / "Materials_and_Methods_GBE_draft.docx"


def set_run_font(style, name: str, size: float, *, bold: bool | None = None) -> None:
    font = style.font
    font.name = name
    font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    style.element.rPr.rFonts.set(qn("w:ascii"), name)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_line_numbering(section) -> None:
    sect_pr = section._sectPr
    for existing in sect_pr.findall(qn("w:lnNumType")):
        sect_pr.remove(existing)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    line_numbers.set(qn("w:distance"), "360")
    sect_pr.append(line_numbers)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_run_font(normal, "Times New Roman", 12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    set_run_font(title, "Times New Roman", 16, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)

    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = styles[name]
        set_run_font(style, "Times New Roman", size, bold=True)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1

    caption = styles["Caption"]
    set_run_font(caption, "Times New Roman", 10)
    caption.font.italic = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.line_spacing = 1
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)

    if "Author Input" not in styles:
        warning = styles.add_style("Author Input", WD_STYLE_TYPE.PARAGRAPH)
        warning.base_style = normal
    warning = styles["Author Input"]
    set_run_font(warning, "Times New Roman", 11, bold=True)
    warning.font.color.rgb = None
    warning.paragraph_format.left_indent = Inches(0.25)
    warning.paragraph_format.right_indent = Inches(0.25)
    warning.paragraph_format.space_before = Pt(6)
    warning.paragraph_format.space_after = Pt(6)
    warning.paragraph_format.line_spacing = 1.15


def build_reference() -> None:
    document = Document()
    configure_styles(document)
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        add_line_numbering(section)
        add_page_number(section.footer.paragraphs[0])
    document.core_properties.title = "Materials and Methods — GBE working draft"
    document.core_properties.subject = "Desmognathus TE and microscopy methods"
    document.core_properties.author = "Desmognathus TE project"
    document.save(REFERENCE)


def postprocess(path: Path) -> None:
    document = Document(path)
    configure_styles(document)
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        add_line_numbering(section)
        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.clear()
        add_page_number(paragraph)

    for paragraph in document.paragraphs:
        if "AUTHOR INPUT REQUIRED" in paragraph.text:
            paragraph.style = document.styles["Author Input"]

    document.core_properties.title = "Materials and Methods — GBE working draft"
    document.core_properties.subject = "Evidence-backed working draft with embedded supporting figures"
    document.core_properties.author = "Desmognathus TE project"
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    build_reference()
    subprocess.run(
        [
            "pandoc",
            str(source),
            "--from=markdown+link_attributes",
            "--to=docx",
            f"--reference-doc={REFERENCE}",
            f"--resource-path={source.parent}:{ROOT / 'Publication'}:{ROOT}",
            f"--output={output}",
        ],
        cwd=ROOT,
        check=True,
    )
    postprocess(output)
    print(output)


if __name__ == "__main__":
    main()
