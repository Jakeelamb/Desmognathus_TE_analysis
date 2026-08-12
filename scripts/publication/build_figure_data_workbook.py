"""Build the single researcher-facing workbook of exact ggplot figure inputs."""

from __future__ import annotations

import argparse
import hashlib
import math
import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[2]
PUBLICATION = ROOT / "Publication"
DATASETS = PUBLICATION / "datasets"
MANIFEST_PATH = DATASETS / "DATASET_MANIFEST.csv"
COLUMN_INVENTORY_PATH = DATASETS / "COLUMN_INVENTORY.csv"
DEFAULT_OUTPUT = PUBLICATION / "Desmognathus_figure_data.xlsx"

FIXED_DOCUMENT_TIME = datetime.fromisoformat("1980-01-01T00:00:00")
INVALID_EXCEL_TEXT = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


@dataclass(frozen=True)
class FigureInput:
    supplement_id: str
    sheet_name: str
    figure: str
    panel: str
    figure_use: str


FIGURE_INPUTS = (
    FigureInput("S02", "S02_TE_Panel", "Figure 1", "B", "Counts the 34-species genomic TE panel."),
    FigureInput(
        "S03",
        "S03_Path_Panel",
        "Figure 1",
        "B",
        "Counts the 24-species phenotype/path panel and its TE34 overlap.",
    ),
    FigureInput(
        "S04",
        "S04_Path_Traits",
        "Figure 1",
        "A",
        "Supplies the species labels and morphology traits displayed beside the Path24 phylogeny.",
    ),
    FigureInput(
        "S05",
        "S05_Assembly_QC",
        "Figure 1",
        "C",
        "Plots assembly span against contig N50; labels the largest spans and highest N50.",
    ),
    FigureInput(
        "S06",
        "S06_dnaPipeTE_QC",
        "Figure 1",
        "D",
        "Plots repeat-aligned fraction against unresolved order mass and marks retry runs.",
    ),
    FigureInput(
        "S08",
        "S08_TE_Diversity",
        "Figure 2",
        "A",
        "Filters the classified-only table to te_level=order; plots Shannon entropy and Gini-Simpson across 34 species. Unresolved aligned-base mass remains separate S07 QC and is not a diversity category.",
    ),
    FigureInput(
        "S09",
        "S09_PCA_Scores",
        "Figure 2",
        "B",
        "Filters to superfamily/classified_conditional; plots PC1 versus PC2 and labels the ten most distant scores.",
    ),
    FigureInput(
        "S10",
        "S10_PCA_Variance",
        "Figure 2",
        "B-C",
        "Supplies PC1/PC2 axis variance and plots the first ten scree values.",
    ),
    FigureInput(
        "S12",
        "S12_Repeat_Landscape",
        "Figure 3",
        "A-B",
        "Groups the six largest repeat orders plus Other, zero-fills bins, and summarizes young repeats below 5% divergence.",
    ),
    FigureInput(
        "S14",
        "S14_LTR_Elements",
        "Figure 4",
        "A",
        "Filters primary_ltr_element=TRUE; plots all 380 retained element ratios.",
    ),
    FigureInput(
        "S15",
        "S15_LTR_Species",
        "Figure 4",
        "A-B",
        "Filters to primary_iqr_filtered_coverage_ge_80pct with positive medians; plots species medians and intervals.",
    ),
    FigureInput(
        "S16",
        "S16_LTR_Coverage",
        "Figure 4",
        "B",
        "Supplies the exact retained-element support count for each species.",
    ),
    FigureInput(
        "S19",
        "S19_Morph_Species",
        "Figures 5 and 6",
        "5A-B; 6A-C",
        "Plots species cell/nucleus estimates and supplies morphology values joined into the pairwise panels.",
    ),
    FigureInput(
        "S22",
        "S22_IOD_Species",
        "Figure 6",
        "A-B",
        "Plots relative nuclear-IOD estimates and conditional intervals.",
    ),
    FigureInput(
        "S24",
        "S24_IOD_QC",
        "Figure 6",
        "A-B",
        "Marks species that fail the prespecified IOD quality-balance thresholds.",
    ),
    FigureInput(
        "S26",
        "S26_Pairwise_PGLS",
        "Figure 6",
        "A-C",
        "Supplies the frozen Pagel-lambda sensitivity lines, intervals, and annotations.",
    ),
    FigureInput(
        "S30",
        "S30_Path_Models",
        "Figure 7",
        "A",
        "Filters analysis_type=primary, ranks by delta_CICc, and plots the first ten models.",
    ),
    FigureInput(
        "S36",
        "S36_Path_Sim",
        "Figure 7",
        "B",
        "Plots true-class top-model recovery by generating class and residual regime.",
    ),
    FigureInput(
        "S40",
        "S40_TE_IOD_Overlap",
        "Figure S1",
        "A",
        "Plots the exact 21-species TE-diversity and relative-IOD overlap without imputation.",
    ),
)

SHEET_ORDER = (
    "README",
    "Figure_Index",
    "Column_Dictionary",
    *(item.sheet_name for item in FIGURE_INPUTS),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _native_value(value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"Workbook cannot store non-finite value: {value}")
    if isinstance(value, str):
        return INVALID_EXCEL_TEXT.sub("", value)
    return value


def _literal_text(cell: Cell, value: str) -> None:
    cell.value = value
    if value.startswith("="):
        cell.data_type = "s"


def _style_header(cell: Cell) -> None:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _format_for_column(column: str, value: Any) -> str:
    lower = column.lower()
    if isinstance(value, bool):
        return "General"
    if isinstance(value, int):
        return "#,##0"
    if not isinstance(value, float):
        return "General"
    if "p_value" in lower or lower.endswith("_p") or lower == "p":
        return "0.00000"
    if any(token in lower for token in ("fraction", "proportion", "rate", "weight", "r2")):
        return "0.0000"
    if any(token in lower for token in ("count", "_bp", "_n", "n_", "length")):
        return "#,##0.000"
    return "0.000000"


def _set_widths(sheet, rows: list[list[Any]], cap: int = 46) -> None:
    for index in range(1, sheet.max_column + 1):
        values = [row[index - 1] for row in rows if index <= len(row)]
        width = max((len(str(value)) for value in values if value is not None), default=8)
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(
            max(width + 2, 10), cap
        )


def _add_table_sheet(workbook: Workbook, name: str, frame: pd.DataFrame) -> None:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = True
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = None
    sheet.row_dimensions[1].height = 22

    rows: list[list[Any]] = [list(frame.columns)]
    rows.extend(
        [
            [_native_value(value) for value in row]
            for row in frame.itertuples(index=False, name=None)
        ]
    )
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column_index)
            if isinstance(value, str):
                _literal_text(cell, value)
            else:
                cell.value = value
            if row_index == 1:
                _style_header(cell)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                cell.number_format = _format_for_column(str(rows[0][column_index - 1]), value)

    _set_widths(sheet, rows)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.auto_filter.ref = sheet.dimensions


def _add_readme(workbook: Workbook, manifest_sha: str) -> None:
    sheet = workbook.create_sheet("README")
    sheet.sheet_view.showGridLines = True
    sheet.freeze_panes = "A2"
    rows = [
        ("item", "details"),
        ("Workbook", "Desmognathus figure data workbook"),
        (
            "Purpose",
            "Exact controlled non-tree CSV inputs consumed by the current ggplot figure builder. Source tables are copied without row filtering or re-analysis.",
        ),
        (
            "Authority",
            "Generated convenience view only; it is not a second data authority. Publication/datasets/DATASET_MANIFEST.csv and its SHA-256 hashes remain authoritative.",
        ),
        (
            "Supporting data",
            "Tree files, tree edge/node tables, and other supporting or audit supplements remain under Publication/datasets in their appropriate standard formats and are intentionally not duplicated here.",
        ),
        (
            "Navigation",
            "Use Figure_Index for the exact sheet, panel, filter, and transformation; use Column_Dictionary for source field definitions.",
        ),
        ("Direct figure-input tables", str(len(FIGURE_INPUTS))),
        ("Rebuild command", "make figure-data"),
        ("Figure source", "scripts/publication/build_gbe_figures.R"),
        ("Dataset-manifest SHA-256", manifest_sha),
        (
            "Relative IOD boundary",
            "Image-derived relative phenotype; not an independently calibrated absolute genome size or C-value.",
        ),
        (
            "LTR boundary",
            "The LTR statistic is a mapping/deletion-footprint proxy; not a direct ectopic-recombination, solo-LTR, or DNA-loss rate.",
        ),
        (
            "Path-analysis boundary",
            "PGLS and path-model outputs are exploratory sensitivity analyses and do not establish a unique causal direction.",
        ),
        (
            "Tree provenance boundary",
            "Supplied by Alex Pyron; publication/archive identifier, calibration method, and branch-length provenance remain unresolved.",
        ),
        (
            "Join boundary",
            "Genomic and microscopy measurements are joined at the species level and do not represent measurements from the same animals.",
        ),
    ]
    for row_index, row_values in enumerate(rows, start=1):
        for column_index, value in enumerate(row_values, start=1):
            cell = sheet.cell(row_index, column_index, value)
            if row_index == 1:
                cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet["B6"].hyperlink = "#'Figure_Index'!A1"
    sheet.auto_filter.ref = f"A1:B{len(rows)}"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 115
    for row_index in range(2, len(rows) + 1):
        sheet.row_dimensions[row_index].height = 30
    sheet.page_setup.orientation = "portrait"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1


def _load_release() -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    manifest = pd.read_csv(MANIFEST_PATH, low_memory=False)
    manifest_rows = manifest.set_index("supplement_id", drop=False)
    expected_ids = {item.supplement_id for item in FIGURE_INPUTS}
    if not expected_ids.issubset(set(manifest_rows.index)):
        raise ValueError("Figure workbook input is missing from DATASET_MANIFEST.csv")

    frames: dict[str, pd.DataFrame] = {}
    for item in FIGURE_INPUTS:
        row = manifest_rows.loc[item.supplement_id]
        source = DATASETS / str(row["filename"])
        if not source.is_file():
            raise FileNotFoundError(source)
        if sha256(source) != str(row["output_sha256"]):
            raise ValueError(f"Manifest hash mismatch for {item.supplement_id}: {source}")
        frame = pd.read_csv(source, low_memory=False)
        if frame.shape != (int(row["rows"]), int(row["columns"])):
            raise ValueError(f"Manifest shape mismatch for {item.supplement_id}: {frame.shape}")
        frames[item.supplement_id] = frame
    return manifest, frames


def _add_figure_index(workbook: Workbook, manifest: pd.DataFrame) -> None:
    manifest_rows = manifest.set_index("supplement_id", drop=False)
    rows: list[dict[str, Any]] = []
    for item in FIGURE_INPUTS:
        source = manifest_rows.loc[item.supplement_id]
        rows.append(
            {
                "figure": item.figure,
                "panel": item.panel,
                "sheet": item.sheet_name,
                "supplement_id": item.supplement_id,
                "dataset_title": source["title"],
                "publication_file": f"datasets/{source['filename']}",
                "canonical_source": source["source_path"],
                "rows": int(source["rows"]),
                "columns": int(source["columns"]),
                "release_status": source["release_status"],
                "figure_use_or_filter": item.figure_use,
                "dataset_description": source["description"],
            }
        )
    frame = pd.DataFrame(rows)
    _add_table_sheet(workbook, "Figure_Index", frame)
    sheet = workbook["Figure_Index"]
    for row_number in range(2, sheet.max_row + 1):
        sheet_cell = sheet.cell(row_number, 3)
        sheet_cell.hyperlink = f"#'{sheet_cell.value}'!A1"
        publication_cell = sheet.cell(row_number, 6)
        publication_cell.hyperlink = str(publication_cell.value)
        canonical_cell = sheet.cell(row_number, 7)
        canonical_cell.hyperlink = f"../{canonical_cell.value}"


def _add_column_dictionary(workbook: Workbook) -> None:
    inventory = pd.read_csv(COLUMN_INVENTORY_PATH, low_memory=False)
    sheet_for_id = {item.supplement_id: item.sheet_name for item in FIGURE_INPUTS}
    inventory = inventory[inventory["supplement_id"].isin(sheet_for_id)].copy()
    inventory.insert(1, "sheet", inventory["supplement_id"].map(sheet_for_id))
    expected_columns = sum(
        len(pd.read_csv(DATASETS / filename, nrows=0).columns)
        for filename in inventory["filename"].unique()
    )
    if len(inventory) != expected_columns:
        raise ValueError("Column inventory does not cover every figure-input field exactly once")
    _add_table_sheet(workbook, "Column_Dictionary", inventory)
    sheet = workbook["Column_Dictionary"]
    for row_number in range(2, sheet.max_row + 1):
        cell = sheet.cell(row_number, 2)
        cell.hyperlink = f"#'{cell.value}'!A1"


def build_workbook() -> Workbook:
    manifest, frames = _load_release()
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Desmognathus figure data"
    workbook.properties.subject = "Exact controlled data inputs for the current ggplot figures"
    workbook.properties.creator = "Desmognathus study team"
    workbook.properties.keywords = (
        "Desmognathus, figure data, ggplot, transposable elements, morphology"
    )
    workbook.properties.description = (
        "Generated convenience workbook; manifest-bound CSV files remain authoritative."
    )
    workbook.properties.created = FIXED_DOCUMENT_TIME
    workbook.properties.modified = FIXED_DOCUMENT_TIME
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False

    _add_readme(workbook, sha256(MANIFEST_PATH))
    _add_figure_index(workbook, manifest)
    _add_column_dictionary(workbook)
    for item in FIGURE_INPUTS:
        _add_table_sheet(workbook, item.sheet_name, frames[item.supplement_id])

    if tuple(workbook.sheetnames) != SHEET_ORDER:
        raise AssertionError("Workbook sheet order drifted")
    return workbook


def _normalize_xlsx_archive(source: Path, destination: Path) -> None:
    with (
        zipfile.ZipFile(source, "r") as source_zip,
        zipfile.ZipFile(
            destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as destination_zip,
    ):
        for name in sorted(source_zip.namelist()):
            payload = source_zip.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    b'<dcterms:modified xsi:type="dcterms:W3CDTF">1980-01-01T00:00:00Z</dcterms:modified>',
                    payload,
                )
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            destination_zip.writestr(
                info, payload, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9
            )


def verify_workbook(path: Path) -> None:
    manifest, frames = _load_release()
    workbook = load_workbook(path, read_only=False, data_only=False)
    if tuple(workbook.sheetnames) != SHEET_ORDER:
        raise ValueError("Workbook does not contain the expected ordered sheets")

    manifest_rows = manifest.set_index("supplement_id", drop=False)
    for item in FIGURE_INPUTS:
        sheet = workbook[item.sheet_name]
        frame = frames[item.supplement_id]
        if (sheet.max_row, sheet.max_column) != (len(frame) + 1, len(frame.columns)):
            raise ValueError(f"Workbook shape mismatch in {item.sheet_name}")
        if [cell.value for cell in sheet[1]] != list(frame.columns):
            raise ValueError(f"Workbook header mismatch in {item.sheet_name}")
        if len(sheet.tables) != 0 or sheet.freeze_panes != "B2":
            raise ValueError(f"Workbook navigation/style contract missing in {item.sheet_name}")
        if sha256(DATASETS / str(manifest_rows.loc[item.supplement_id, "filename"])) != str(
            manifest_rows.loc[item.supplement_id, "output_sha256"]
        ):
            raise ValueError(f"Source changed while verifying {item.supplement_id}")

    formula_errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError(f"Unexpected formula in {sheet.title}!{cell.coordinate}")
                if isinstance(cell.value, str) and cell.value in formula_errors:
                    raise ValueError(f"Spreadsheet error in {sheet.title}!{cell.coordinate}")
    workbook.close()


def write_workbook(output: Path) -> bool:
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    lock_path = output.with_name(f".~lock.{output.name}#")
    if output.exists() and lock_path.exists():
        raise RuntimeError(
            f"Workbook is open and cannot be replaced safely: {output}. "
            "Close it in LibreOffice/Excel and rerun the build."
        )
    with tempfile.TemporaryDirectory(prefix="desmognathus-workbook-") as temporary:
        temporary_path = Path(temporary)
        raw = temporary_path / "raw.xlsx"
        normalized = temporary_path / "normalized.xlsx"
        workbook = build_workbook()
        workbook.save(raw)
        _normalize_xlsx_archive(raw, normalized)
        verify_workbook(normalized)
        changed = not output.exists() or sha256(output) != sha256(normalized)
        if changed:
            output.write_bytes(normalized.read_bytes())
    verify_workbook(output)
    return changed


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    changed = write_workbook(arguments.output)
    output = arguments.output.resolve()
    state = "wrote" if changed else "verified unchanged"
    print(
        f"{state}: {output} | sheets={len(SHEET_ORDER)} | "
        f"figure_tables={len(FIGURE_INPUTS)} | sha256={sha256(output)}"
    )


if __name__ == "__main__":
    main()
