"""Build the researcher-facing Materials and Methods DOCX from Markdown."""

from __future__ import annotations

import argparse
import html
import re
import tempfile
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = ROOT / "Publication/methods/MATERIALS_AND_METHODS_RESEARCHER_DRAFT.md"
DEFAULT_OUTPUT = ROOT / "Publication/methods/MATERIALS_AND_METHODS_RESEARCHER_DRAFT.docx"
FIXED_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
FIXED_DATETIME = datetime(2026, 8, 13, tzinfo=UTC)
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
INLINE_PATTERN = re.compile(r"(`[^`]+`|\*[^*]+\*|<sub>.*?</sub>|<sup>.*?</sup>)")
GENERATOR_COMMENT = "Generated from MATERIALS_AND_METHODS_RESEARCHER_DRAFT.md"
MARKDOWN_LINK_PATTERN = re.compile(r"\[([^]]+)]\([^)]+\)")


def set_run_font(run, *, size: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic


def set_keep_with_next(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    properties.append(OxmlElement("w:keepNext"))


def set_widow_control(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    properties.append(OxmlElement("w:widowControl"))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.first_line_indent = Inches(0.25)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15

    if "Manuscript Title" not in styles:
        title = styles.add_style("Manuscript Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title = styles["Manuscript Title"]
    title.base_style = normal
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(15)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Inches(0)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    heading = styles["Heading 1"]
    heading.font.name = FONT
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = BLACK
    heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.paragraph_format.first_line_indent = Inches(0)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True


def add_inline_text(paragraph, text: str) -> None:
    for token in INLINE_PATTERN.split(text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=10.5)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=11, italic=True)
        elif token.startswith("<sub>") and token.endswith("</sub>"):
            run = paragraph.add_run(token[5:-6])
            set_run_font(run, size=9)
            run.font.subscript = True
        elif token.startswith("<sup>") and token.endswith("</sup>"):
            run = paragraph.add_run(token[5:-6])
            set_run_font(run, size=9)
            run.font.superscript = True
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=11)


def parse_markdown(source: Path) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append(("paragraph", " ".join(paragraph_lines).strip()))
            paragraph_lines.clear()

    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("<!--") or line.endswith("-->"):
            continue
        if not line:
            flush_paragraph()
        elif line.startswith("# "):
            flush_paragraph()
            blocks.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            flush_paragraph()
            blocks.append(("heading", line[3:].strip()))
        elif line.startswith("### "):
            flush_paragraph()
            blocks.append(("subheading", line[4:].strip()))
        else:
            paragraph_lines.append(line)
    flush_paragraph()
    return blocks


def build_document(source: Path, output: Path, *, overwrite: bool = False) -> None:
    if output.is_file() and not overwrite:
        raise RuntimeError(
            "Refusing to overwrite an author-edited Word document; "
            "synchronize its prose intentionally before rebuilding"
        )

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    configure_styles(document)

    for kind, text in parse_markdown(source):
        if kind == "title":
            paragraph = document.add_paragraph(style="Manuscript Title")
            run = paragraph.add_run(text)
            set_run_font(run, size=15, bold=True)
        elif kind in {"heading", "subheading"}:
            paragraph = document.add_paragraph(
                style="Heading 1" if kind == "heading" else "Heading 2"
            )
            run = paragraph.add_run(text)
            set_run_font(run, size=11.5, bold=True)
            set_keep_with_next(paragraph)
        else:
            paragraph = document.add_paragraph(style="Normal")
            add_inline_text(paragraph, text)
            set_widow_control(paragraph)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    add_page_number(footer_paragraph)

    properties = document.core_properties
    properties.title = "Materials and Methods"
    properties.subject = "Researcher-facing draft for the Desmognathus TE study"
    properties.author = "Desmognathus TE Study"
    properties.keywords = "Desmognathus; transposable elements; microscopy; genome size"
    properties.comments = GENERATOR_COMMENT
    properties.created = FIXED_DATETIME
    properties.modified = FIXED_DATETIME
    properties.last_modified_by = "Desmognathus TE Study"
    properties.revision = 1

    output.parent.mkdir(parents=True, exist_ok=True)
    lock_path = output.with_name(f".~lock.{output.name}#")
    if lock_path.exists():
        raise RuntimeError(f"Refusing to overwrite an open LibreOffice document: {lock_path}")

    with tempfile.TemporaryDirectory(prefix="desmog-methods-") as temp_dir:
        raw_path = Path(temp_dir) / "raw.docx"
        document.save(raw_path)
        with (
            zipfile.ZipFile(raw_path, "r") as archive_in,
            zipfile.ZipFile(
                output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
            ) as archive_out,
        ):
            for name in sorted(archive_in.namelist()):
                info = zipfile.ZipInfo(name, FIXED_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                archive_out.writestr(
                    info, archive_in.read(name), compress_type=zipfile.ZIP_DEFLATED
                )


def markdown_visible_text(text: str) -> str:
    """Return the visible text represented by the supported Markdown subset."""
    text = MARKDOWN_LINK_PATTERN.sub(r"\1", text)
    text = re.sub(r"</?(?:sub|sup|u)>", "", text)
    text = text.replace("`", "").replace("*", "")
    text = re.sub(r"\\([\\`*{}\[\]()#+\-.!_|>])", r"\1", text)
    return html.unescape(text)


def check_document_matches_source(source: Path, output: Path) -> None:
    expected = [(kind, markdown_visible_text(text)) for kind, text in parse_markdown(source)]
    document = Document(output)
    actual: list[tuple[str, str]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name == "Manuscript Title":
            kind = "title"
        elif paragraph.style.name == "Heading 2":
            kind = "subheading"
        elif paragraph.style.name.startswith("Heading"):
            kind = "heading"
        else:
            kind = "paragraph"
        actual.append((kind, text))
    if actual != expected:
        raise RuntimeError("Author-edited Word document and Markdown companion are out of sync")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--check",
        action="store_true",
        help="Verify visible DOCX prose matches the Markdown companion",
    )
    parser.add_argument(
        "--replace",
        action="store_true",
        help="Explicitly replace an existing DOCX from Markdown",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    if args.check:
        if not output.is_file():
            raise SystemExit(f"Missing generated document: {output}")
        check_document_matches_source(source, output)
        print(f"PASS: visible prose in {output} matches {source}")
        return
    build_document(source, output, overwrite=args.replace)
    print(f"Wrote {output}")


if __name__ == "__main__":
    main()
