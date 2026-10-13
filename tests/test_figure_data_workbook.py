from __future__ import annotations

import math
import re
import runpy
import subprocess
import sys
import tomllib
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIGURE_BUILDER = ROOT / "scripts/publication/build_gbe_figures.R"
WORKBOOK_BUILDER = ROOT / "scripts/publication/build_figure_data_workbook.py"
WORKBOOK_MODULE = runpy.run_path(str(WORKBOOK_BUILDER))
DATASETS = WORKBOOK_MODULE["DATASETS"]
DEFAULT_OUTPUT = WORKBOOK_MODULE["DEFAULT_OUTPUT"]
FIGURE_INPUTS = WORKBOOK_MODULE["FIGURE_INPUTS"]
MANIFEST_PATH = WORKBOOK_MODULE["MANIFEST_PATH"]
SHEET_ORDER = WORKBOOK_MODULE["SHEET_ORDER"]
_native_value = WORKBOOK_MODULE["_native_value"]
sha256 = WORKBOOK_MODULE["sha256"]
verify_workbook = WORKBOOK_MODULE["verify_workbook"]
write_workbook = WORKBOOK_MODULE["write_workbook"]


def _values_equal(actual: object, expected: object) -> bool:
    if actual is None or expected is None:
        return actual is expected
    if isinstance(actual, bool) or isinstance(expected, bool):
        return actual is expected
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return math.isclose(float(actual), float(expected), rel_tol=1e-12, abs_tol=1e-12)
    return actual == expected


def _workbook_under_test(tmp_path: Path) -> Path:
    lock = DEFAULT_OUTPUT.with_name(f".~lock.{DEFAULT_OUTPUT.name}#")
    if not lock.exists():
        return DEFAULT_OUTPUT
    rebuilt = tmp_path / DEFAULT_OUTPUT.name
    subprocess.run(
        [
            sys.executable,
            str(WORKBOOK_BUILDER),
            "--output",
            str(rebuilt),
        ],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return rebuilt


def test_open_canonical_workbook_uses_a_temporary_build(tmp_path: Path, monkeypatch) -> None:
    opened = tmp_path / "editor" / DEFAULT_OUTPUT.name
    opened.parent.mkdir()
    opened.write_bytes(b"open workbook")
    lock = opened.with_name(f".~lock.{opened.name}#")
    lock.write_text("open")
    monkeypatch.setitem(globals(), "DEFAULT_OUTPUT", opened)

    rebuilt = _workbook_under_test(tmp_path)

    assert rebuilt == tmp_path / opened.name
    verify_workbook(rebuilt)
    assert opened.read_bytes() == b"open workbook"
    assert lock.read_text() == "open"


def test_workbook_scope_matches_non_tree_ggplot_inputs() -> None:
    manifest = pd.read_csv(MANIFEST_PATH).set_index("supplement_id")
    expected_files = {
        str(manifest.loc[item.supplement_id, "filename"])
        for item in FIGURE_INPUTS
        if item.input_role in {"direct_plot_input", "validation_input"}
    }
    figure_source = FIGURE_BUILDER.read_text(encoding="utf-8")
    direct_files = set(re.findall(r'read_supp\("([^"]+\.csv)"\)', figure_source))

    tree_files = {
        "Supplementary_Data_S38_phylogeny_edges.csv",
        "Supplementary_Data_S39_phylogeny_nodes.csv",
    }
    assert direct_files - tree_files == expected_files
    assert '"desmognathus_time_tree_path24_v1.nwk"' in figure_source
    assert len(FIGURE_INPUTS) == 16
    assert all(item.supplement_id not in {"S38", "S39"} for item in FIGURE_INPUTS)
    assert {
        item.supplement_id for item in FIGURE_INPUTS if item.input_role == "calibration_proof"
    } == {"S22"}


def test_diversity_sheet_is_the_wang_style_top10_release(tmp_path: Path) -> None:
    manifest = pd.read_csv(MANIFEST_PATH).set_index("supplement_id")
    source = pd.read_csv(DATASETS / str(manifest.loc["S08", "filename"]))
    expected_columns = [
        "species",
        "n_superfamilies_included",
        "shannon_entropy",
        "gini_simpson",
    ]
    assert source.shape == (34, 4)
    assert list(source.columns) == expected_columns
    assert source["species"].is_unique
    assert set(source["n_superfamilies_included"]) == {10}

    workbook = load_workbook(_workbook_under_test(tmp_path), read_only=True, data_only=True)
    sheet = workbook["S08_TE_Diversity"]
    assert (sheet.max_row, sheet.max_column) == (35, 4)
    assert [cell.value for cell in sheet[1]] == expected_columns
    workbook.close()

    diversity_input = next(item for item in FIGURE_INPUTS if item.supplement_id == "S08")
    assert "ten most abundant" in diversity_input.figure_use
    assert "reclosed" in diversity_input.figure_use
    assert "not a diversity category" in diversity_input.figure_use
    assert "te_level=order" not in diversity_input.figure_use


def test_bonsai_figure_sheets_bind_exact_annotations_and_validation(
    tmp_path: Path,
) -> None:
    workbook = load_workbook(_workbook_under_test(tmp_path), read_only=True, data_only=True)
    tips = workbook["S42_Bonsai_Tips"]
    validation = workbook["S43_Bonsai_Validation"]

    assert (tips.max_row, tips.max_column) == (35, 6)
    tip_rows = list(tips.iter_rows(min_row=2, values_only=True))
    strategy_index = [cell.value for cell in tips[1]].index("reproductive_strategy")
    strategies = [row[strategy_index] for row in tip_rows]
    assert strategies.count("aquatic_eggs") == 7
    assert strategies.count("direct_development") == 3
    assert strategies.count("terrestrial_eggs_aquatic_larvae") == 24
    assert strategies.count("unassigned") == 0

    assert (validation.max_row, validation.max_column) == (4, 15)
    headers = [cell.value for cell in validation[1]]
    rows = [dict(zip(headers, row, strict=True)) for row in validation.iter_rows(
        min_row=2, values_only=True
    )]
    by_strategy = {row["strategy"]: row for row in rows}
    assert by_strategy["aquatic_eggs"]["feature_bootstrap_edge_delimited_count"] == 55
    assert by_strategy["direct_development"]["feature_bootstrap_edge_delimited_count"] == 70
    assert by_strategy["terrestrial_eggs_aquatic_larvae"][
        "feature_bootstrap_edge_delimited_count"
    ] == 2
    assert {row["global_fitch_observed_steps"] for row in rows} == {2}
    workbook.close()


def test_workbook_defaults_to_genome_size_and_keeps_iod_as_calibration_proof(
    tmp_path: Path,
) -> None:
    workbook = load_workbook(_workbook_under_test(tmp_path), read_only=True, data_only=True)
    assert "S22_IOD_Calibration" in workbook.sheetnames
    assert "S40_TE_Genome_Overlap" not in workbook.sheetnames
    assert "S40_TE_IOD_Overlap" not in workbook.sheetnames

    path_headers = [cell.value for cell in workbook["S04_Path_Traits"][1]]
    pairwise = workbook["S26_Pairwise_PGLS"]
    calibration_headers = [cell.value for cell in workbook["S22_IOD_Calibration"][1]]

    assert "genome_size_pg_fuscus_anchored" in path_headers
    assert "relative_iod_index" not in path_headers
    assert "relative_iod_index" in calibration_headers
    assert "genome_size_pg_fuscus_anchored" in calibration_headers

    pairwise_rows = list(pairwise.iter_rows(values_only=True))
    serialized = "\n".join(",".join(str(value) for value in row) for row in pairwise_rows)
    assert "genome_size_vs_nucleus_area" in serialized
    assert "genome_size_vs_cell_area" in serialized
    assert "relative_iod" not in serialized
    workbook.close()


def test_canonical_workbook_has_exact_source_values_and_navigation(tmp_path: Path) -> None:
    workbook_path = _workbook_under_test(tmp_path)
    verify_workbook(workbook_path)
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    assert tuple(workbook.sheetnames) == SHEET_ORDER
    assert len(workbook.sheetnames) == 19
    assert {"Tree_Newick", "S38_Tree_Edges", "S39_Tree_Nodes"}.isdisjoint(workbook.sheetnames)
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert len(set(workbook.sheetnames)) == len(workbook.sheetnames)

    manifest = pd.read_csv(MANIFEST_PATH).set_index("supplement_id")
    for item in FIGURE_INPUTS:
        source_path = DATASETS / str(manifest.loc[item.supplement_id, "filename"])
        source = pd.read_csv(source_path, low_memory=False)
        sheet = workbook[item.sheet_name]
        assert sheet.freeze_panes == "B2"
        assert sheet.sheet_view.showGridLines is True
        assert len(sheet.tables) == 0
        assert sheet.auto_filter.ref == sheet.dimensions
        assert [cell.value for cell in sheet[1]] == list(source.columns)

        actual_rows = sheet.iter_rows(min_row=2, values_only=True)
        expected_rows = (
            tuple(_native_value(value) for value in row)
            for row in source.itertuples(index=False, name=None)
        )
        for row_number, (actual, expected) in enumerate(
            zip(actual_rows, expected_rows, strict=True), start=2
        ):
            assert all(
                _values_equal(actual_value, expected_value)
                for actual_value, expected_value in zip(actual, expected, strict=True)
            ), f"source parity failed in {item.sheet_name} row {row_number}"
    workbook.close()


def test_workbook_uses_plain_spreadsheet_formatting(tmp_path: Path) -> None:
    workbook = load_workbook(_workbook_under_test(tmp_path), read_only=False, data_only=False)
    for sheet in workbook.worksheets:
        assert sheet.sheet_view.showGridLines is True
        assert sheet.sheet_properties.tabColor is None
        assert not sheet.merged_cells.ranges
        assert len(sheet.tables) == 0
        for row_number, row in enumerate(sheet.iter_rows(), start=1):
            for cell in row:
                if cell.value is None:
                    continue
                assert cell.fill.fill_type is None
                assert cell.border.left.style is None
                assert cell.border.right.style is None
                assert cell.border.top.style is None
                assert cell.border.bottom.style is None
                assert cell.font.color is None or cell.font.color.type == "theme"
                assert cell.font.underline is None
                assert cell.font.bold is (row_number == 1)
    workbook.close()


def test_open_workbook_is_never_overwritten(tmp_path: Path) -> None:
    output = tmp_path / "open.xlsx"
    output.write_bytes(b"original")
    lock = tmp_path / ".~lock.open.xlsx#"
    lock.write_text("open")
    try:
        write_workbook(output)
    except RuntimeError as error:
        assert "Workbook is open" in str(error)
    else:
        raise AssertionError("expected a live workbook lock to block replacement")
    assert output.read_bytes() == b"original"


def test_workbook_is_a_deterministic_manifest_bound_build(tmp_path: Path) -> None:
    rebuilt = tmp_path / "Desmognathus_figure_data.xlsx"
    assert write_workbook(rebuilt)
    with zipfile.ZipFile(rebuilt) as archive:
        core = ElementTree.fromstring(archive.read("docProps/core.xml"))
    modified = core.find("{http://purl.org/dc/terms/}modified")
    assert modified is not None
    assert modified.text == "1980-01-01T00:00:00Z"
    assert modified.attrib["{http://www.w3.org/2001/XMLSchema-instance}type"] == "dcterms:W3CDTF"
    if not DEFAULT_OUTPUT.with_name(f".~lock.{DEFAULT_OUTPUT.name}#").exists():
        assert sha256(rebuilt) == sha256(DEFAULT_OUTPUT)
    assert not write_workbook(rebuilt)


def test_workbook_exposes_boundaries_and_make_target(tmp_path: Path) -> None:
    workbook = load_workbook(_workbook_under_test(tmp_path), read_only=True, data_only=True)
    readme = "\n".join(
        str(cell.value)
        for row in workbook["README"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    index_rows = list(workbook["Figure_Index"].iter_rows(values_only=True))
    workbook.close()

    for phrase in [
        "not a second data authority",
        "conditional D. fuscus assembly-anchored pg/1C estimate",
        "S22 preserves the underlying IOD measurement and conversion proof",
        "mapping/deletion-footprint proxy",
        "do not establish a unique causal direction",
        "publication/archive identifier",
        "make figure-data",
    ]:
        assert phrase in readme
    assert len(index_rows) == 17  # header + 16 non-tree figure/validation inputs

    manifest = tomllib.loads((ROOT / "pixi.toml").read_text(encoding="utf-8"))
    task = manifest["feature"]["analysis"]["tasks"]["figure-data"]
    assert "validate" in task["depends-on"]
    assert "build_figure_data_workbook.py" in task["cmd"]
    dry_run = subprocess.run(
        ["make", "--dry-run", "figure-data"],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    assert "--locked --no-install figure-data" in dry_run.stdout
