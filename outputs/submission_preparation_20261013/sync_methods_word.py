"""Synchronize reviewed prose while preserving unchanged Word package parts."""

from copy import deepcopy
from difflib import SequenceMatcher
from io import BytesIO
import json
from pathlib import Path
import re
from runpy import run_path
import sys
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
RUN = Path(__file__).resolve().parent
builder = run_path(ROOT / "scripts/publication/build_methods_document.py")
source = ROOT / "Publication/methods/MATERIALS_AND_METHODS_RESEARCHER_DRAFT.md"
output = source.with_suffix(".docx")
baseline_source = (
    Path(sys.argv[1]).resolve() if len(sys.argv) == 2
    else RUN / "baseline/Publication/methods" / source.name
)
baseline_output = baseline_source.with_suffix(".docx")

if output.with_name(f".~lock.{output.name}#").exists():
    raise RuntimeError("Word draft has a LibreOffice lock")
if output.read_bytes() != baseline_output.read_bytes():
    raise RuntimeError("Word draft changed after baseline; reconcile author edits first")
builder["check_document_matches_source"](baseline_source, baseline_output)
old_blocks = builder["parse_markdown"](baseline_source)
new_blocks = builder["parse_markdown"](source)
document = Document(baseline_output)
paragraphs = list(document.paragraphs)
if len(paragraphs) != len(old_blocks):
    raise RuntimeError("Unexpected baseline paragraph structure")


def make_paragraph(kind, text):
    style = {
        "title": "Manuscript Title",
        "heading": "Heading 1",
        "subheading": "Heading 2",
        "paragraph": "Normal",
    }[kind]
    paragraph = document.add_paragraph(style=style)
    paragraph._p.getparent().remove(paragraph._p)
    if kind == "paragraph":
        # Preserve markup for italic/subscript/superscript runs, while removing
        # Markdown escapes that have no visible counterpart in Word.
        text = re.sub(r"\\([\\`*{}\[\]()#+\-.!_|>])", r"\1", text)
        builder["add_inline_text"](paragraph, text)
        builder["set_widow_control"](paragraph)
    else:
        run = paragraph.add_run(builder["markdown_visible_text"](text))
        builder["set_run_font"](run, size=15 if kind == "title" else 11.5, bold=True)
        builder["set_keep_with_next"](paragraph)
    return paragraph._p


new_nodes = []
unchanged = 0
for operation, old_start, old_stop, new_start, new_stop in SequenceMatcher(
    a=old_blocks, b=new_blocks, autojunk=False
).get_opcodes():
    if operation == "equal":
        new_nodes.extend(deepcopy(p._p) for p in paragraphs[old_start:old_stop])
        unchanged += old_stop - old_start
    else:
        new_nodes.extend(make_paragraph(*block) for block in new_blocks[new_start:new_stop])

body = document._element.body
section_properties = deepcopy(body.find(qn("w:sectPr")))
for child in list(body):
    body.remove(child)
for node in new_nodes:
    body.append(node)
body.append(section_properties)

serialized = BytesIO()
document.save(serialized)
with ZipFile(BytesIO(serialized.getvalue())) as changed:
    new_xml = changed.read("word/document.xml")

result = BytesIO()
with ZipFile(baseline_output) as original, ZipFile(result, "w") as updated:
    for info in original.infolist():
        updated.writestr(
            info,
            new_xml if info.filename == "word/document.xml" else original.read(info.filename),
        )

pending = RUN / "Methods.pending.docx"
pending.write_bytes(result.getvalue())
builder["check_document_matches_source"](source, pending)
with ZipFile(baseline_output) as original, ZipFile(pending) as updated:
    changed_parts = [
        name for name in original.namelist() if original.read(name) != updated.read(name)
    ]
    assert changed_parts == ["word/document.xml"], changed_parts
pending.replace(output)
receipt = {
    "status": "PASS",
    "baseline_paragraphs": len(old_blocks),
    "baseline_source": str(baseline_source),
    "final_paragraphs": len(new_blocks),
    "unchanged_paragraphs_preserved": unchanged,
    "changed_package_parts": changed_parts,
    "styles_headers_footers_and_page_settings": "preserved byte-for-byte",
}
(RUN / "word_sync.json").write_text(json.dumps(receipt, indent=2) + "\n")
print(json.dumps(receipt, indent=2))
