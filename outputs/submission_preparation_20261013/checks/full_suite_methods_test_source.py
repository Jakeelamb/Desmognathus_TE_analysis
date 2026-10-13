from __future__ import annotations

from pathlib import Path
from runpy import run_path
from xml.etree import ElementTree
from zipfile import ZipFile

import pandas as pd
import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
METHODS = ROOT / "Publication/methods"


def _docx_text(path: Path) -> str:
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(document_xml)
    word_text = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    return " ".join((node.text or "") for node in root.iter(word_text))


def test_methods_evidence_matrix_uses_only_active_sources() -> None:
    evidence = pd.read_csv(METHODS / "METHODS_EVIDENCE_MATRIX.csv")
    assert evidence["section_id"].tolist() == [
        f"M{number:02d}" for number in range(1, len(evidence) + 1)
    ]
    assert set(f"M{number:02d}" for number in range(1, 27)) <= set(evidence["section_id"])
    assert not evidence["primary_sources"].isna().any()

    forbidden = ("input_data/", "path_analysis/", "results/", "plans/")
    for source_list in evidence["primary_sources"].astype(str):
        for source in (item.strip() for item in source_list.split(";")):
            assert not source.startswith(forbidden), source
            if source == "none" or source.startswith("doi:"):
                continue
            assert (ROOT / source).is_file(), source


def test_methods_scaffold_records_finalized_release_boundaries() -> None:
    scaffold = (METHODS / "METHODS_BULLET_SCAFFOLD.md").read_text(encoding="utf-8")
    required_phrases = {
        "one 34-row, four-column table in TE34 species order",
        "ten most abundant strictly positive superfamilies were selected independently",
        "descending-abundance ties were resolved deterministically by feature name",
        "ten retained proportions were reclosed to one",
        "`n_superfamilies_included` is ten for every species",
        "Natural-log Shannon entropy was `H′ = −Σpᵢ ln(pᵢ)`",
        "explicitly named Gini-Simpson index was `1 − Σpᵢ²`",
        "These are the only paper-facing diversity indices",
        "do not document whether those ten were reclosed",
        "Diversity uses ten species-specific superfamilies, whereas the Bonsai representation and retained PCA audit use the 24 superfamilies positive in all 34 species",
        "upstream commit `b16ef6f0a164610327022c7d0b92de82b3596aad`",
        "2 observed Fitch steps versus a 9.4132-step permutation-null mean",
        "55/200 (27.5%) and 70/200 (35.0%)",
        "not a species phylogeny",
        "not a biological-specimen bootstrap",
        "Unresolved aligned-base mass remains only in S06-S07 dnaPipeTE quality-control tables",
        "left LTR, right LTR, and internal region",
        "inclusive, within-species, two-sided Tukey filter",
        "380 primary elements",
        "whitespace-delimited `DOMAIN|MODEL` annotations",
        "0.12 µm/pixel",
        "Seven species are represented by one animal",
        "`nuclear_iod_by_image.csv`",
        "fail this gate and remain visibly flagged",
        "Default analysis trait: `target genome size (pg/1C)",
        "Express the path node as **Genome size**",
        "S21-S25 preserve the underlying IOD measurement, quality, and calibration audit",
        "data/te_genome_size_overlap.csv",
    }
    assert all(phrase in scaffold for phrase in required_phrases)
    assert "Hill q1" not in scaffold
    assert "Hill q2" not in scaffold
    assert "Pielou" not in scaffold
    assert "mass-aware" not in scaffold
    assert "mass_aware_unresolved_bin" not in scaffold
    assert "136 rows" not in scaffold
    assert "four exact 34-species strata" not in scaffold
    assert "68-row, five-column" not in scaffold
    assert "order-level stratum is the primary" not in scaffold
    assert "superfamily-level stratum is a sensitivity" not in scaffold
    assert "Observed richness" not in scaffold


def test_s08_metadata_declares_canonical_top10_superfamily_diversity() -> None:
    metadata = run_path(ROOT / "scripts/publication_metadata.py")
    dependencies = metadata["SUPPLEMENT_DEPENDENCIES"]
    column_definitions = metadata["SUPPLEMENT_COLUMN_DEFINITIONS"]
    supplement_metadata = metadata["SUPPLEMENT_METADATA"]

    assert dependencies["S08"] == (
        "analyses/01_transposable_elements/data/te_superfamily_composition.csv",
        "analyses/01_transposable_elements/recompute_pca.py",
    )
    assert set(column_definitions["S08"]) == {
        "n_superfamilies_included",
        "shannon_entropy",
        "gini_simpson",
    }

    description = supplement_metadata["S08"]["description"]
    manuscript_use = supplement_metadata["S08"]["manuscript_use"]
    assert "34-row, four-column Wang-style TE-diversity table" in description
    assert "ten most abundant positive classified superfamilies" in description
    assert "reclosed to one" in description
    assert "species-specific top ten" in manuscript_use
    assert "Wang et al. did not document" in manuscript_use
    assert "Unresolved aligned-base mass remains only" in manuscript_use
    assert "S06-S07" in manuscript_use

    assert dependencies["S41"] == (
        "analyses/01_transposable_elements/data/wang2025_figure2_te_diversity.csv",
        "analyses/01_transposable_elements/data/wang2025_figure2_te_diversity.provenance.json",
        "analyses/01_transposable_elements/build_cross_study_diversity.py",
        "analyses/01_transposable_elements/data/te_diversity.csv",
        "analyses/02_morphology/data/relative_nuclear_iod_species.csv",
        "data/identity/te34_panel.csv",
        "data/identity/path24_panel.csv",
    )
    assert dependencies["S42"] == (
        "data/identity/te34_panel.csv",
        "analyses/03_phylogenetic_path/data/organismal_traits.csv",
        "analyses/01_transposable_elements/bonsai_comparison/validate_strategy_localization.py",
        "analyses/01_transposable_elements/bonsai_comparison/data/te34_bonsai_strategy_validation_manifest.json",
    )
    assert dependencies["S43"] == (
        "analyses/01_transposable_elements/bonsai_comparison/data/te34_bonsai_strategy_label_permutation_tests.csv",
        "analyses/01_transposable_elements/bonsai_comparison/data/te34_bonsai_feature_bootstrap_strategy_stability.csv",
        "analyses/01_transposable_elements/bonsai_comparison/validate_strategy_localization.py",
        "analyses/01_transposable_elements/bonsai_comparison/data/te34_bonsai_strategy_validation_manifest.json",
    )
    assert supplement_metadata["S42"]["release_status"] == "main_candidate"
    assert supplement_metadata["S43"]["release_status"] == "main_candidate"
    assert "not a species phylogeny" in supplement_metadata["S42"]["manuscript_use"]
    assert "biological-specimen bootstrap" in supplement_metadata["S43"]["manuscript_use"]


def test_conditional_assembly_anchor_metadata_is_traceable_and_limited() -> None:
    metadata = run_path(ROOT / "scripts/publication_metadata.py")
    dependencies = metadata["SUPPLEMENT_DEPENDENCIES"]
    column_definitions = metadata["SUPPLEMENT_COLUMN_DEFINITIONS"]
    supplement_metadata = metadata["SUPPLEMENT_METADATA"]

    anchor_dependencies = {
        "analyses/02_morphology/data/relative_nuclear_iod_species.csv",
        "analyses/02_morphology/provenance/iod_source_manifest.json",
        "analyses/02_morphology/validate_release.py",
    }
    assert anchor_dependencies <= set(dependencies["S01"])
    assert anchor_dependencies - {
        "analyses/02_morphology/data/relative_nuclear_iod_species.csv"
    } <= set(dependencies["S22"])

    for supplement_id in ["S01", "S22"]:
        description = supplement_metadata[supplement_id]["description"]
        manuscript_use = supplement_metadata[supplement_id]["manuscript_use"]
        assert "conditional" in description.casefold()
        assert "16.1-Gbp" in description
        assert "0.978 Gbp per pg" in description
        assert "not a measured C-value" in manuscript_use

    s01_definitions = column_definitions["S01"]
    assert "Exact S22 projection" in s01_definitions["step06_genome_size_pg_fuscus_anchored"]
    assert "not a measured C-value" in s01_definitions["step06_genome_size_pg_fuscus_anchored"]

    s22_definitions = column_definitions["S22"]
    expected_calibration_columns = {
        "iod_ratio_to_fuscus",
        "genome_size_pg_fuscus_anchored",
        "genome_size_pg_ci_low",
        "genome_size_pg_ci_high",
        "genome_size_reference_species",
        "genome_size_reference_pg",
        "genome_size_calibration_pg_per_iod",
        "genome_size_rank",
    }
    assert expected_calibration_columns <= set(s22_definitions)
    assert "16.1 Gbp / 0.978 Gbp per pg" in s22_definitions["genome_size_reference_pg"]
    assert "anchor-value uncertainty is excluded" in s22_definitions["genome_size_pg_ci_low"]
    assert "not a measured C-value" in s22_definitions["genome_size_pg_fuscus_anchored"]


def test_genome_size_is_the_default_analysis_trait_and_iod_remains_provenance() -> None:
    metadata = run_path(ROOT / "scripts/publication_metadata.py")
    dependencies = metadata["SUPPLEMENT_DEPENDENCIES"]
    definitions = metadata["SUPPLEMENT_COLUMN_DEFINITIONS"]
    supplement_metadata = metadata["SUPPLEMENT_METADATA"]

    analysis_supplements = ["S04", *[f"S{number:02d}" for number in range(26, 38)], "S40"]
    for supplement_id in analysis_supplements:
        fields = " ".join(supplement_metadata[supplement_id].values()).casefold()
        assert "genome size" in fields or "genome-size" in fields, supplement_id
        assert "relative iod is primary" not in fields, supplement_id
        assert "relative-iod analysis" not in fields, supplement_id

        dependency_set = set(dependencies[supplement_id])
        assert "analyses/02_morphology/data/relative_nuclear_iod_species.csv" in dependency_set
        assert "analyses/02_morphology/provenance/iod_source_manifest.json" in dependency_set

    s04_columns = {
        "genome_size_pg_fuscus_anchored",
        "genome_size_pg_ci_low",
        "genome_size_pg_ci_high",
        "n_genome_size_nuclei",
        "n_genome_size_images",
        "n_genome_size_specimens",
        "genome_size_definition",
    }
    assert s04_columns <= set(definitions["S04"])
    assert not any(
        "relative_iod" in column or column.startswith("n_iod_") for column in definitions["S04"]
    )
    assert "genome_size_log10_z" in definitions["S27"]
    assert {
        "genome_size_pg_fuscus_anchored",
        "genome_size_pg_ci_low",
        "genome_size_pg_ci_high",
        "genome_size_definition",
    } <= set(definitions["S40"])
    assert not {
        "relative_iod_index",
        "relative_iod_ci_low",
        "relative_iod_ci_high",
        "iod_definition",
    } & set(definitions["S40"])

    for supplement_id in ("S21", "S22", "S23", "S24", "S25"):
        fields = " ".join(supplement_metadata[supplement_id].values()).casefold()
        assert "iod" in fields, supplement_id
    assert "default downstream genome-size trait" in supplement_metadata["S22"]["manuscript_use"]


def test_analysis_tables_use_genome_size_terms_not_relative_iod_terms() -> None:
    expected_columns = {
        "species",
        "tree_tip",
        "genome_size_pg_fuscus_anchored",
        "genome_size_pg_ci_low",
        "genome_size_pg_ci_high",
        "nucleus_area_um2",
        "cell_area_um2",
        "n_genome_size_nuclei",
        "n_genome_size_images",
        "n_genome_size_specimens",
        "n_size_cells",
        "n_size_images",
        "genome_size_definition",
    }
    path_traits = pd.read_csv(ROOT / "analyses/03_phylogenetic_path/data/path24_traits.csv")
    assert set(path_traits.columns) == expected_columns
    assert len(path_traits) == 24

    pairwise = pd.read_csv(ROOT / "analyses/03_phylogenetic_path/data/pairwise_pgls.csv")
    assert set(pairwise["comparison"]) == {
        "genome_size_vs_nucleus_area",
        "genome_size_vs_cell_area",
        "nucleus_area_vs_cell_area",
    }
    assert "genome_size_pg_fuscus_anchored" in set(pairwise["left_metric"])

    standardized = pd.read_csv(
        ROOT / "analyses/03_phylogenetic_path/data/path_standardized_traits.csv"
    )
    assert set(standardized.columns) == {
        "species",
        "tree_tip",
        "genome_size_log10_z",
        "nucleus_area_log10_z",
        "cell_area_log10_z",
        "genome_size_definition",
    }

    regressions = pd.read_csv(
        ROOT / "analyses/03_phylogenetic_path/data/pairwise_phylogenetic_regressions.csv"
    )
    assert set(regressions["relation_id"]) == {
        "genome_to_nucleus",
        "nucleus_to_genome",
        "genome_to_cell",
        "cell_to_genome",
        "nucleus_to_cell",
        "cell_to_nucleus",
        "genome_and_nucleus_to_cell",
        "genome_and_cell_to_nucleus",
        "nucleus_and_cell_to_genome",
    }
    assert set(regressions["outcome"]) == {"genome_size", "nucleus_size", "cell_size"}
    assert set(regressions["predictor"]) == {"genome_size", "nucleus_size", "cell_size"}

    definition = (
        "conditional_fuscus_assembly_anchored_genome_size_pg_1c; iod_derived_not_measured_c_value"
    )
    for name in (
        "pairwise_phylogenetic_regressions.csv",
        "evolutionary_model_sensitivity.csv",
        "path_model_comparison.csv",
        "path_edge_estimates.csv",
        "path_model_stability.csv",
        "path_tree_sensitivity.csv",
        "path_measurement_sensitivity.csv",
        "path_leave_one_species_out.csv",
        "path_simulation_calibration.csv",
        "path_dsep_tests.csv",
    ):
        frame = pd.read_csv(ROOT / "analyses/03_phylogenetic_path/data" / name)
        assert set(frame["path_trait_definition"]) == {definition}, name
        serialized = frame.to_csv(index=False).casefold()
        assert "relative_iod" not in serialized, name
        assert "iod_bridge" not in serialized, name
        assert "iod_collider" not in serialized, name
        assert "iod_nucleus_only" not in serialized, name
        assert "iod_cell_only" not in serialized, name

    overlap = pd.read_csv(ROOT / "analyses/03_phylogenetic_path/data/te_genome_size_overlap.csv")
    assert len(overlap) == 21
    assert {
        "genome_size_pg_fuscus_anchored",
        "genome_size_pg_ci_low",
        "genome_size_pg_ci_high",
        "genome_size_definition",
    } <= set(overlap)
    assert not {
        "relative_iod_index",
        "relative_iod_ci_low",
        "relative_iod_ci_high",
        "iod_definition",
    } & set(overlap)


def test_active_docs_do_not_restore_relative_iod_as_the_analysis_default() -> None:
    documents = (
        ROOT / "README.md",
        ROOT / "docs/REPRODUCIBILITY.md",
        ROOT / "docs/CLEANUP_LEDGER.md",
        ROOT / "analyses/02_morphology/README.md",
        ROOT / "analyses/03_phylogenetic_path/README.md",
        ROOT / "Publication/README.md",
        ROOT / "Publication/datasets/README.md",
        METHODS / "README.md",
        METHODS / "METHODS_BULLET_SCAFFOLD.md",
        METHODS / "METHODS_EVIDENCE_MATRIX.csv",
        ROOT / "Publication/figures/README.md",
        ROOT / "Publication/figures/GBE_FIGURE_PLAN.md",
        ROOT / "Publication/figures/FIGURE_LEGENDS_AND_ALT_TEXT.md",
    )
    stale_claims = (
        "relative iod remains the primary",
        "relative iod is primary",
        "continues to use relative iod as its primary",
        "primary variable is relative image-derived nuclear iod",
        "express the path node as relative nuclear iod",
        "model contains only relative nuclear iod",
        "s04, s26, and s27-s37 express the path24 image phenotype as relative nuclear iod",
    )
    for path in documents:
        text = path.read_text(encoding="utf-8").casefold()
        assert not any(claim in text for claim in stale_claims), path

    combined = "\n".join(path.read_text(encoding="utf-8") for path in documents)
    assert "conditional *D. fuscus* assembly-anchored genome size" in combined
    assert "S21-S25" in combined
    assert "not an independent genome-size assay" in combined


def test_researcher_methods_docx_is_plain_current_and_singular() -> None:
    source = METHODS / "MATERIALS_AND_METHODS_RESEARCHER_DRAFT.md"
    document = METHODS / "MATERIALS_AND_METHODS_RESEARCHER_DRAFT.docx"
    assert source.is_file()
    assert document.is_file()
    assert sorted(path.name for path in METHODS.glob("*.docx")) == [document.name]
    assert not (METHODS / "MATERIALS_AND_METHODS_NARRATIVE_DRAFT.docx").exists()
    assert not (METHODS / "METHODS_BULLET_SCAFFOLD_WORD_READY.docx").exists()

    forbidden = (
        "[verified]",
        "[boundary]",
        "[author input]",
        "[author confirmed]",
        "[historical command]",
        "[planned",
        "author-facing factual outline",
        "analyses not to rerun",
        "citation needed",
        "citations needed",
        "will be inserted",
        "will be added",
        "will be cited",
        "[confirm",
        "/home/",
        "/nfs/",
        "mass-aware",
        "mass aware",
        "mass_aware_unresolved_bin",
        "136 rows",
        "four exact 34-species strata",
        "values that included unresolved repeat mass",
        "order-level shannon and gini-simpson values were the primary diversity measures",
        "all positive classified categories contributed to the diversity calculations",
        "observed richness was reported as a supporting count",
    )
    narrative = _docx_text(document)
    folded = narrative.casefold()
    assert not any(phrase.casefold() in folded for phrase in forbidden)
    assert "Materials and Methods" in narrative
    assert "Unresolved mass did not enter the diversity calculations" in narrative
    assert "ten most abundant strictly positive TE superfamilies" in narrative
    assert "do not state whether their retained top-ten abundances were reclosed" in narrative
    assert "leaving 24 shared superfamilies" in narrative
    assert "380 elements across all 30 species" in narrative
    assert "specimens 32469 and 32470" in narrative
    assert "16.462167689 pg/1C" in narrative
    assert "Pagel-lambda phylogenetic generalized least squares" in narrative
    assert "25 directed acyclic graphs" in narrative
    assert "does not validate the absolute genome-size scale" in narrative
    assert "exact 21 species with complete TE, cellular, and maximum-SVL data" in narrative
    assert "expanded body-role comparison comprised eight distinct equivalence classes" in narrative
    assert "animal-aware body-adjusted sensitivity" in narrative
    assert "allometry refits re-estimated lambda" in narrative
    assert "brightfield microscopy" in narrative
    assert "40× objective" in narrative
    assert "independently validated C-values" in narrative
    assert "TE composition and maximum snout-vent length have not been incorporated" not in narrative
    assert "60×" not in narrative
    assert "All animals were ethanol-fixed" not in narrative
    assert "SRA Toolkit" not in narrative
    assert "fastp" not in narrative
    assert "DeepTE" not in narrative
    assert "TEClass2" not in narrative

    builder = run_path(ROOT / "scripts/publication/build_methods_document.py")
    builder["check_document_matches_source"](source, document)


def test_methods_builder_refuses_to_overwrite_author_edits(tmp_path: Path) -> None:
    builder = run_path(ROOT / "scripts/publication/build_methods_document.py")
    source = tmp_path / "methods.md"
    source.write_text("# Materials and Methods\n\nOriginal source.\n", encoding="utf-8")
    output = tmp_path / "methods.docx"
    builder["build_document"](source, output)

    document = Document(output)
    document.paragraphs[-1].text = "Author-edited prose."
    document.save(output)
    source.write_text("# Materials and Methods\n\nUpdated source.\n", encoding="utf-8")

    with pytest.raises(RuntimeError, match="author-edited Word document"):
        builder["build_document"](source, output)

    with pytest.raises(RuntimeError, match="out of sync"):
        builder["check_document_matches_source"](source, output)


def test_methods_builder_preserves_subsection_heading_text(tmp_path: Path) -> None:
    builder = run_path(ROOT / "scripts/publication/build_methods_document.py")
    source = tmp_path / "methods.md"
    source.write_text(
        "# Materials and Methods\n\n## Comparative analyses\n\n"
        "### Paired measurements\n\nMeasurements remain paired.\n",
        encoding="utf-8",
    )
    output = tmp_path / "methods.docx"
    builder["build_document"](source, output)
    builder["check_document_matches_source"](source, output)
    document = Document(output)
    heading = document.paragraphs[2]
    assert heading.text == "Paired measurements"
    assert heading.style.name == "Heading 2"
    assert "###" not in _docx_text(output)


def test_methods_readme_separates_internal_evidence_from_shareable_prose() -> None:
    readme = (METHODS / "README.md").read_text(encoding="utf-8")
    normalized = " ".join(readme.split())
    assert "internal ledger" in normalized
    assert "author-edited Word draft" in normalized
    assert "refuses to replace an existing author-edited DOCX" in normalized
    assert "make methods-word" in normalized
    assert "no audit tags" in normalized
    assert "two synchronized Word" not in normalized.casefold()


def test_unresolved_methods_are_recorded_without_recreating_old_paths() -> None:
    required = pd.read_csv(METHODS / "METHODS_REQUIRED_AUTHOR_INPUT.csv").set_index("item_id")
    assert required.loc["W12", "status"] == "partial"
    assert required.loc["W16", "status"] == "partial"
    assert "multimapper policy" in required.loc["W12", "required_detail"]
    assert "checkpoint-to-mask-release mapping" in required.loc["W16", "required_detail"]
    assert required.loc["W15", "status"] == "completed_exploratory_documented_author_review_pending"
    assert required.loc["W06", "status"] == "partial"
    assert "not standards 413/414" in required.loc["W06", "current_repository_evidence"]
    assert "run map" in required.loc["W06", "current_repository_evidence"]

    old_image_table = ROOT / "analyses/02_morphology/data/relative_nuclear_iod_by_image.csv"
    new_image_table = ROOT / "analyses/02_morphology/data/nuclear_iod_by_image.csv"
    assert not old_image_table.exists()
    assert new_image_table.is_file()
